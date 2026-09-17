import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generar_documentacion_exhaustiva(path_docx):
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Styles Setup
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    NAVY = RGBColor(0x04, 0x0D, 0x1A)
    BLUE = RGBColor(0x25, 0x63, 0xEB)
    SLATE_DARK = RGBColor(0x0F, 0x17, 0x2A)
    SLATE_GRAY = RGBColor(0x47, 0x55, 0x69)

    # Title Banner
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    run_sub = p_title.add_run("DEPARTAMENTO DE SISTEMAS — GRUPO HUERTA\n")
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = BLUE

    run_main = p_title.add_run("MANUAL Y DOCUMENTACIÓN EXHAUSTIVA DE INFRAESTRUCTURA DE PORTALES\n")
    run_main.font.size = Pt(18)
    run_main.font.bold = True
    run_main.font.color.rgb = NAVY

    run_desc = p_title.add_run("Configuración de cPanel, GitHub Secrets, Versiones PHP, Variables .env y Arquitectura de Software\n")
    run_desc.font.size = Pt(11)
    run_desc.font.italic = True
    run_desc.font.color.rgb = SLATE_GRAY

    doc.add_paragraph()

    def add_custom_heading(text, level=1):
        h = doc.add_paragraph()
        run = h.add_run(text)
        run.font.bold = True
        if level == 1:
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(8)
            run.font.size = Pt(14)
            run.font.color.rgb = BLUE
        elif level == 2:
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            run.font.size = Pt(12)
            run.font.color.rgb = SLATE_DARK
        elif level == 3:
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            run.font.size = Pt(11)
            run.font.color.rgb = BLUE
        return h

    def add_callout(text, title="📌 NOTA TÉCNICA Y DE SEGURIDAD", fill_hex="EFF6FF", border_hex="2563EB"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, fill_hex)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        
        p = cell.paragraphs[0]
        run_t = p.add_run(f"{title}\n")
        run_t.bold = True
        run_t.font.color.rgb = BLUE
        run_t.font.size = Pt(10.5)

        run_b = p.add_run(text)
        run_b.font.size = Pt(10)
        run_b.font.color.rgb = SLATE_DARK
        doc.add_paragraph()

    # SECTION 1: ARQUITECTURA GENERAL
    add_custom_heading("1. Resumen Ejecutivo y Entornos Duales", level=1)
    doc.add_paragraph(
        "El proyecto implementa la infraestructura tecnológica centralizada para la Dirección de Sistemas de Grupo Huerta. "
        "Consta de dos portales web interconectados mediante una arquitectura descentralizada (Zero-Storage Central), lo que garantiza "
        "que cada sucursal conserve la soberanía total de sus bases de datos en su propio cPanel."
    )

    table_env = doc.add_table(rows=3, cols=4)
    table_env.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Portal", "Dominio / Subdominio", "Rama Git", "Directorio Local en PC"]
    for i, h in enumerate(headers):
        cell = table_env.cell(0, i)
        set_cell_background(cell, "0F172A")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)

    data_env = [
        ["Portal Agencia (Divolavilla)", "portal.divolavilla.com", "main", "C:\\Users\\Sistemas\\Documents\\Systems Portal"],
        ["Portal Maestro (Grupo Huerta)", "portal.grupohuerta.mx", "gh", "C:\\Users\\Sistemas\\Documents\\Systems Portal_GH"]
    ]

    for row_idx, row_data in enumerate(data_env, start=1):
        for col_idx, text in enumerate(row_data):
            cell = table_env.cell(row_idx, col_idx)
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.5)

    doc.add_paragraph()

    # SECTION 2: SECRETS DE GITHUB Y CI/CD
    add_custom_heading("2. Configuración de Variables Secretas en GitHub (GitHub Secrets)", level=1)
    doc.add_paragraph(
        "Para lograr el despliegue automático continuo (CI/CD) sin exponer contraseñas en el código fuente público o privado, "
        "se configuraron las Variables de Secretos en el repositorio de GitHub:"
    )
    doc.add_paragraph("📍 Ruta en GitHub: Repositorio ➔ Settings ➔ Secrets and variables ➔ Actions ➔ Repository secrets")

    table_sec = doc.add_table(rows=5, cols=3)
    table_sec.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_sec = ["Nombre de la Variable (Secret)", "Descripción / Propósito", "Ejemplo de Valor"]
    for i, h in enumerate(headers_sec):
        cell = table_sec.cell(0, i)
        set_cell_background(cell, "2563EB")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)

    data_sec = [
        ["FTP_SERVER", "Dirección IP o Host del servidor cPanel", "srv01.vwserver.com.mx (o IP del servidor)"],
        ["FTP_USERNAME", "Usuario de acceso FTP o cPanel", "grupohuerta / divolavilla"],
        ["FTP_PASSWORD", "Contraseña segura del usuario de cPanel", "••••••••••••••••"],
        ["FTP_SERVER_DIR", "Directorio destino remoto en el servidor", "public_html/sistemas/"]
    ]

    for row_idx, row_data in enumerate(data_sec, start=1):
        for col_idx, text in enumerate(row_data):
            cell = table_sec.cell(row_idx, col_idx)
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.5)

    doc.add_paragraph()
    doc.add_paragraph(
        "Funcionamiento de GitHub Actions (.github/workflows/deploy_gh.yml):\n"
        "Cada vez que se ejecuta un comando `git push origin gh` o `git push origin main`, GitHub activa un flujo de trabajo que toma "
        "los archivos actualizados y los sube vía FTP seguro a la carpeta `public_html/sistemas/` del cPanel de manera 100% automatizada."
    )

    # SECTION 3: CONFIGURACIÓN CPANEL Y MYSQL
    add_custom_heading("3. Configuración Detallada de cPanel y Base de Datos MySQL", level=1)
    
    add_custom_heading("3.1 Creación de Base de Datos y Usuarios MySQL", level=2)
    doc.add_paragraph("1. Iniciar sesión en cPanel de Grupo Huerta (o Divolavilla).")
    doc.add_paragraph("2. Ir a la sección Bases de Datos ➔ Bases de datos MySQL.")
    doc.add_paragraph("3. En Crear una nueva base de datos, escribir: grupohue_sistemas.")
    doc.add_paragraph("4. En Usuarios MySQL ➔ Añadir nuevo usuario, crear el usuario: grupohue_admin con contraseña segura.")
    doc.add_paragraph("5. En Añadir usuario a la base de datos (sección inferior):")
    doc.add_paragraph("   • Seleccionar Usuario: grupohue_admin y Base de datos: grupohue_sistemas.")
    doc.add_paragraph("   • Hacer clic en Añadir.")
    doc.add_paragraph("   • Marcar la casilla TODOS LOS PRIVILEGIOS (ALL PRIVILEGES) y presionar Hacer Cambios.")

    add_custom_heading("3.2 Importación de Tablas en phpMyAdmin", level=2)
    doc.add_paragraph("1. Entra a cPanel ➔ phpMyAdmin.")
    doc.add_paragraph("2. Seleccionar la base de datos grupohue_sistemas en la columna izquierda.")
    doc.add_paragraph("3. Hacer clic en la pestaña Importar (Import).")
    doc.add_paragraph("4. Seleccionar el archivo SQL database/schema_cpanel.sql y hacer clic en Ejecutar.")

    add_custom_heading("3.3 Selección de Versión de PHP en MultiPHP Manager (Solución Error 500)", level=2)
    doc.add_paragraph("1. Buscar la herramienta Administrador de MultiPHP (MultiPHP Manager) en cPanel.")
    doc.add_paragraph("2. Buscar en la lista el subdominio portal.grupohuerta.mx y marcar su casilla.")
    doc.add_paragraph("3. En el desplegable Versión de PHP de la derecha, seleccionar PHP 8.1 o PHP 8.2.")
    doc.add_paragraph("4. Presionar Aplicar. (Esto regenera la directiva ea-php81 en Apache y resuelve bloqueos 500).")

    add_custom_heading("3.4 Permisos de Archivos y Carpetas (suPHP / FastCGI)", level=2)
    doc.add_paragraph("En servidores cPanel con arquitectura suPHP, se debe cumplir estrictamente:")
    doc.add_paragraph("• Archivos .php y .htaccess ➔ Permisos 0644 (Lectura/Escritura para el propietario, Lectura para el grupo).")
    doc.add_paragraph("• Carpetas y Directorios ➔ Permisos 0755.")

    add_callout(
        "Si un archivo .php tiene permisos 0666 o 0777 en cPanel, el módulo de seguridad suPHP rehúsa ejecutarlo y arroja un HTTP Error 500 automático por razones de seguridad. Mantener siempre en 0644.",
        title="⚠️ REGLA CRÍTICA DE PERMISOS CPANEL"
    )

    # SECTION 4: VARIABLES LOCALES .ENV Y CONEXION.PHP
    add_custom_heading("4. Variables de Entorno Locales (.env) y conexion.php", level=1)
    
    add_custom_heading("4.1 Estructura del Archivo .env Local", level=2)
    doc.add_paragraph("En la raíz del proyecto local en tu computadora se cuenta con el archivo protegido .env:")
    
    p_env = doc.add_paragraph(
        "DB_SERVER=192.168.26.5\n"
        "DB_NAME=gedas\n"
        "DB_USER=sa\n"
        "DB_PASS=sa\n"
        "CPANEL_URL=https://portal.grupohuerta.mx/recibir_reporte.php\n"
        "TOKEN_SECRETO=GedasDivolavilla2026!"
    )
    p_env.paragraph_format.left_indent = Inches(0.5)

    add_custom_heading("4.2 Manejo de Conexión PDO en conexion.php", level=2)
    doc.add_paragraph(
        "El archivo conexion.php utiliza PDO con soporte para Unix Sockets ('localhost') en cPanel "
        "y maneja excepciones de forma silenciosa para permitir el funcionamiento del portal incluso si la base de datos se encuentra en mantenimiento:"
    )

    p_pdo = doc.add_paragraph(
        "<?php\n"
        "$db_host = getenv('MYSQL_HOST') ?: 'localhost';\n"
        "$db_name = getenv('MYSQL_DB')   ?: 'grupohue_sistemas';\n"
        "$db_user = getenv('MYSQL_USER') ?: 'grupohue_admin';\n"
        "$db_pass = getenv('MYSQL_PASS') ?: 'admin';\n"
        "?>"
    )
    p_pdo.paragraph_format.left_indent = Inches(0.5)

    # SECTION 5: ARQUITECTURA Y INTERFAZ DE USUARIO
    add_custom_heading("5. Flujo de Navegación e Interfaz de Usuario", level=1)
    
    doc.add_paragraph("1. Pantalla de Inicio de Sesión (login.php): Autenticación con bloques try-catch que previenen caídas 500 y ofrecen credenciales maestras de respaldo (admin / Admin123!).")
    doc.add_paragraph("2. Catálogo Maestro de Agencias (menu.php): Presenta las tarjetas de cada sucursal (VW Divol La Villa, Seat La Villa, Cupra Garage) con el botón azul 'Ingresar ➔'.")
    doc.add_paragraph("3. Menú de Módulos Corporativos por Agencia (modulos.php): Presenta los 9 Módulos (Órdenes de Servicio, Inventario de Equipos, Celulares, Licencias, SITE/IDF, Respaldos, Mantenimiento, Correo, Cumplimiento) adaptados a la sucursal seleccionada.")
    doc.add_paragraph("4. Consolidado en Vivo (reportes.php): Petición HTTPS en vivo bajo demanda que consulta los datos sin almacenar registros centralmente.")

    # SECTION 6: EXTRACTOR ENVIAR_REPORTE.PY
    add_custom_heading("6. Extractor e Integrador Python (enviar_reporte.py)", level=1)
    doc.add_paragraph("El script enviar_reporte.py en la PC local de la agencia se encarga de:")
    doc.add_paragraph("1. Conectarse vía pyodbc al servidor SQL Server local (192.168.26.5 / BD gedas).")
    doc.add_paragraph("2. Consultar la tabla SEORDSER para extraer órdenes abiertas, cerradas y canceladas.")
    doc.add_paragraph("3. Empaquetar la información en formato JSON.")
    doc.add_paragraph("4. Enviar una petición HTTPS POST cifrada con el Token de Autorización Bearer GedasDivolavilla2026! hacia https://portal.grupohuerta.mx/recibir_reporte.php.")

    doc.save(path_docx)
    print(f"[ÉXITO] Documentación exhaustiva generada en: {path_docx}")

if __name__ == "__main__":
    path_main = r"C:\Users\Sistemas\Documents\Systems Portal_GH\Documentacion_Proyecto_Portales_Sistemas_GrupoHuerta.docx"
    path_copy = r"C:\Users\Sistemas\Documents\Systems Portal\Documentacion_Proyecto_Portales_Sistemas_GrupoHuerta.docx"
    
    generar_documentacion_exhaustiva(path_main)
    generar_documentacion_exhaustiva(path_copy)
